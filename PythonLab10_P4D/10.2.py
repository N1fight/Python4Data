from docx import Document
from docx.shared import Inches, RGBColor

# Открытие созданного документа
doc = Document("file.docx")

# Добавление изображения
doc.add_paragraph(" ")  # Для отступа
doc.add_picture('image.jpg', width=Inches(4.0))

# Добавление подписи с изменением цвета
caption_paragraph = doc.add_paragraph()
caption_run = caption_paragraph.add_run("Figure 1. Example image")

# Установка цвета текста
caption_run.font.color.rgb = RGBColor(120, 120, 240)  # Установите нужный цвет, например, синий (RGB: 0, 0, 255)

# Сохранение документа
doc.save("file_image.docx")