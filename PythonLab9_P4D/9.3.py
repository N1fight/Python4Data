from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.styles import Alignment

# Открываем созданный файл Excel
wb = load_workbook("Зарплаты.xlsx")
ws = wb.active

# Считываем данные
rows = list(ws.iter_rows(min_row=2, values_only=True))
data = [row for row in rows if row[0] and "Итого" not in row[1]]

# Суммируем зарплаты по отделам
total_salary_by_dept = {}

for row in data:
    dept = row[2]  # Название отдела
    salary = row[5]  # Сумма зарплаты
    if dept not in total_salary_by_dept:
        total_salary_by_dept[dept] = 0
    total_salary_by_dept[dept] += salary

# Вставляем данные для диаграммы в Excel
chart_data = []
for dept, total_salary in total_salary_by_dept.items():
    chart_data.append([dept, total_salary])

# Заголовки для данных диаграммы в K2 и L2
ws.cell(row=2, column=11, value="Отдел").alignment = Alignment(horizontal='center')  # Заголовок "Отдел" в K2
ws.cell(row=2, column=12, value="Сумма зарплаты").alignment = Alignment(horizontal='center')  # Заголовок "Сумма зарплаты" в L2

# Добавляем данные на лист Excel в столбцы K и L, начиная с K3
for idx, row in enumerate(chart_data, start=3):  # Начинаем с 3-й строки
    ws.cell(row=idx, column=11, value=row[0]).alignment = Alignment(horizontal='center')  # Название отдела в K
    ws.cell(row=idx, column=12, value=row[1]).alignment = Alignment(horizontal='center')  # Сумма зарплаты в L

# Устанавливаем ширину столбцов K и L
ws.column_dimensions['K'].width = 20  # Ширина столбца K
ws.column_dimensions['L'].width = 20  # Ширина столбца L

# Создаем круговую диаграмму
chart = PieChart()
data = Reference(ws, min_col=12, min_row=2, max_row=2 + len(chart_data))
categories = Reference(ws, min_col=11, min_row=3, max_row=2 + len(chart_data))
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
chart.title = "Распределение зарплаты по отделам"

# Вставляем диаграмму в свободную область справа от таблицы
chart_position = "N2"  # Позиция диаграммы в N2
ws.add_chart(chart, chart_position)

# Сохраняем файл
wb.save("Зарплаты_с_диаграммой.xlsx")
