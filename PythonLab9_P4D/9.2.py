from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

# Открываем существующий файл
wb = load_workbook("Зарплаты.xlsx")
ws = wb.active

# Установка выравнивания для всех ячеек
for row in ws.iter_rows():
    for cell in row:
        cell.alignment = Alignment(horizontal='center', vertical='center')  # Выравнивание по центру

# Считываем данные
rows = list(ws.iter_rows(min_row=2, values_only=True))
data = [row for row in rows if row[0] and "Итого" not in row[1]]

# Вычисления
max_salary = max(data, key=lambda x: x[5])
min_salary = min(data, key=lambda x: x[5])
average_salary_by_dept = {}

for row in data:
    dept = row[2]
    if dept not in average_salary_by_dept:
        average_salary_by_dept[dept] = []
    average_salary_by_dept[dept].append(row[5])

average_salary_by_dept = {k: sum(v) / len(v) for k, v in average_salary_by_dept.items()}

# Запись результатов в Excel
ws.cell(row=14, column=2, value="Показатель").font = Font(bold=True)
ws.cell(row=14, column=3, value="Значение").font = Font(bold=True)

ws.cell(row=15, column=2, value="Максимальная зарплата:")
ws.cell(row=15, column=3, value=f"{max_salary[1]} - {max_salary[5]} руб.")

ws.cell(row=16, column=2, value="Минимальная зарплата:")
ws.cell(row=16, column=3, value=f"{min_salary[1]} - {min_salary[5]} руб.")

ws.cell(row=17, column=2, value="Средняя зарплата по отделам:").font = Font(bold=True)

row_index = 18  # Начнем с 18 строки
ws.cell(row=row_index, column=2, value="Отдел").font = Font(bold=True)
ws.cell(row=row_index, column=3, value="Средняя зарплата").font = Font(bold=True)

row_index += 1  # Переходим на следующую строку
for dept, avg in average_salary_by_dept.items():
    ws.cell(row=row_index, column=2, value=dept)
    ws.cell(row=row_index, column=3, value=f"{round(avg, 2)} руб.")
    row_index += 1

# Установка ширины столбцов для всего файла
for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter  # Получаем букву столбца
    for cell in column:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = (max_length + 2)  # Добавляем немного пространства
    ws.column_dimensions[column_letter].width = adjusted_width

# Сохраняем изменения в файл
wb.save("Зарплаты.xlsx")
